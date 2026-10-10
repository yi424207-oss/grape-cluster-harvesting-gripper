"""Generate an editable course report and a print-ready Chinese PDF."""
from pathlib import Path
from html import escape
import re
from docx import Document
from docx.shared import Cm, Pt
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os
import shutil
import subprocess

ROOT=Path(__file__).resolve().parent
text=(ROOT/'paper.md').read_text()

def iter_blocks():
    lines=text.splitlines();i=0
    while i<len(lines):
        s=lines[i].strip()
        if not s: i+=1;continue
        if s.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',x) for x in cells):rows.append(cells)
                i+=1
            yield ('table',rows);continue
        if s.startswith('#'):
            yield ('heading',len(s)-len(s.lstrip('#')),s.lstrip('#').strip());i+=1;continue
        yield ('paragraph',s);i+=1

doc=Document()
sec=doc.sections[0];sec.top_margin=Cm(3);sec.bottom_margin=Cm(2.5);sec.left_margin=Cm(3);sec.right_margin=Cm(2.5)
styles=doc.styles
normal=styles['Normal'];normal.font.name='宋体';normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.space_after=Pt(8)
for st,size in [('Title',22),('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
    a=styles[st];a.font.name='黑体';a.font.size=Pt(size);a.font.bold=True;a.font.color.rgb=RGBColor(0,0,0)
    ppr=a.element.get_or_add_pPr()
    bdr=ppr.find(qn('w:pBdr'))
    if bdr is not None:ppr.remove(bdr)
styles['Title'].paragraph_format.space_after=Pt(12)
for st in ('Heading 1','Heading 2','Heading 3'):
    styles[st].paragraph_format.space_before=Pt(16)
    styles[st].paragraph_format.space_after=Pt(8)

def doc_figure(file,caption):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT/file),width=Cm(14.8))
    cap=doc.add_paragraph(caption);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:r.font.size=Pt(10)

for block in iter_blocks():
    kind=block[0]
    if kind=='heading':
        _,level,s=block
        if level==1:doc.add_paragraph(s,style='Title')
        elif level==2:
            doc.add_paragraph(s,style='Heading 1')
            if s=='3 结构与工作原理':doc_figure('assembly_view.png','图1 参数化三维总装闭合位置示意图')
        else:
            doc.add_paragraph(s,style='Heading 2')
            if s=='3.1 总体方案':doc_figure('general_arrangement.png','图2 正视与俯视尺寸方案图 单位 mm')
    elif kind=='table':
        rows=block[1];t=doc.add_table(rows=0,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for row in rows:
            cells=t.add_row().cells
            for c,v in zip(cells,row):c.text=v
        tblpr=t._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','bottom','insideH']:
            elem=OxmlElement('w:'+edge);elem.set(qn('w:val'),'single');elem.set(qn('w:sz'),'8' if edge!='insideH' else '4');borders.append(elem)
        tblpr.append(borders)
    else:
        p=doc.add_paragraph(block[1])
        p.paragraph_format.first_line_indent=Cm(.8)
        if block[1].startswith(('姓名：','指导教师：','关键词：','[','表1','附件A：')):p.paragraph_format.first_line_indent=Cm(0)
doc.save(ROOT/'葡萄果穗柔性采摘手爪课程设计.docx')

office = shutil.which('libreoffice') or shutil.which('soffice')
runtime = os.environ.get('CODEX_PRIMARY_RUNTIME_ROOT')
if not office and runtime:
    candidate=Path(runtime)/'dependencies/native/libreoffice-headless/libreoffice/program/soffice'
    if candidate.exists():office=str(candidate)
if office:
    subprocess.run([office,'-env:UserInstallation=file:///tmp/grape_course_lo',
                    '--headless','--convert-to','pdf','--outdir',str(ROOT),
                    str(ROOT/'葡萄果穗柔性采摘手爪课程设计.docx')],check=True)
    print('Wrote PDF and DOCX')
else:
    print('Wrote DOCX. Install LibreOffice and a Chinese font, then export the DOCX to PDF.')
